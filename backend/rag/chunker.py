"""
Text extraction + chunking for task-scoped RAG.

Reuses backend.tools.file_processing for PDF / text decoding so we share the
same charset detection (UTF-8 → GBK fallback) and PyMuPDF logic as the rest
of the ingest pipeline. Adds a simple sliding-window word chunker — no
tokenizer dependency, predictable RAM footprint on Render free tier.
"""
from __future__ import annotations

import logging
import os
import re
from io import BytesIO
from typing import List

from fastapi import HTTPException

from backend.tools.file_processing import decode_text_bytes, extract_text_from_pdf

logger = logging.getLogger(__name__)


# ─── Limits (per CLAUDE plan) ─────────────────────────────────────────────────
MAX_FILE_BYTES = 64 * 1024 * 1024
MAX_CHUNKS_PER_DOC = 500  # Legacy display constant, no longer a truncation limit.
MAX_CHARS_PER_CHUNK = 2000          # safety belt against runaway docs
DEFAULT_CHUNK_WORDS = 500
DEFAULT_OVERLAP_WORDS = 50


SUPPORTED_EXTS = (".pdf", ".docx", ".pptx", ".md", ".markdown", ".txt", ".rst")


def _guess_kind(filename: str) -> str:
    name = (filename or "").lower()
    for ext in SUPPORTED_EXTS:
        if name.endswith(ext):
            return ext
    return ""


async def extract_text(filename: str, body: bytes) -> str:
    """Decode a KB upload to a single text string.

    Raises HTTPException(400) for unsupported types or decode failure, and
    HTTPException(413) if the file exceeds MAX_FILE_BYTES.
    """
    if len(body) > MAX_FILE_BYTES:
        raise HTTPException(
            status_code=413,
            detail=f"KB file too large ({len(body)} bytes > {MAX_FILE_BYTES}); "
                   f"please split into smaller documents.",
        )

    kind = _guess_kind(filename)
    if not kind:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported KB file type: {os.path.basename(filename)}. "
                   f"Allowed: PDF, DOCX, PPTX, MD, TXT, RST.",
        )

    try:
        if kind == ".pdf":
            text = await extract_text_from_pdf(body)
        elif kind == ".docx":
            from docx import Document

            document = Document(BytesIO(body))
            parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            text = "\n".join(parts)
        elif kind == ".pptx":
            from pptx import Presentation

            presentation = Presentation(BytesIO(body))
            parts = []
            for slide in presentation.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        parts.append(shape.text.strip())
            text = "\n".join(parts)
        else:
            text = await decode_text_bytes(body)
    except Exception as exc:
        logger.warning("Knowledge document extraction failed for %s: %s", filename, type(exc).__name__)
        raise HTTPException(status_code=400, detail="Unable to extract text from this document.") from exc

    text = (text or "").strip()
    if not text:
        raise HTTPException(
            status_code=400,
            detail="KB file appears empty after extraction.",
        )
    return text


def chunk_text(
    text: str,
    *,
    chunk_words: int = DEFAULT_CHUNK_WORDS,
    overlap_words: int = DEFAULT_OVERLAP_WORDS,
) -> List[str]:
    """Compatibility list API, without dropping long tokens or later chunks."""
    return [item["content"] for item in chunk_spans(text, chunk_words=chunk_words, overlap_words=overlap_words)]


def chunk_spans(text: str, *, chunk_words=DEFAULT_CHUNK_WORDS, overlap_words=DEFAULT_OVERLAP_WORDS):
    """Lossless windows with exact character offsets into the original page.

    Prefer word boundaries, but split oversized Chinese/code/formula runs by
    characters. Whitespace and formatting are retained; overlapping windows can
    be reconstructed using offsets without counting their overlap twice.
    """
    if not text or not text.strip():
        return []
    chunk_words = chunk_words if chunk_words > 0 else DEFAULT_CHUNK_WORDS
    overlap_words = overlap_words if 0 <= overlap_words < chunk_words else chunk_words // 10
    result, start = [], 0
    while start < len(text):
        limit = min(len(text), start + MAX_CHARS_PER_CHUNK)
        words = list(re.finditer(r"\S+", text[start:limit]))
        end = start + words[chunk_words].start() if len(words) > chunk_words else limit
        result.append(dict(content=text[start:end], start=start, end=end))
        if end == len(text):
            break
        if len(words) > chunk_words:
            next_start = start + words[chunk_words - overlap_words].start()
        else:
            next_start = end - (min(200, (end - start) // 10) if overlap_words else 0)
        start = max(start + 1, next_start)
    return result
