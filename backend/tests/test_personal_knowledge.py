from io import BytesIO

import fitz
import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from fastapi.testclient import TestClient


def _persist_ready_document(
    *, owner_id: str, filename: str, content: bytes, chunks: list[str]
):
    """Create test knowledge through the quota-managed production boundary."""
    from backend.db.knowledge_repository import get_document, replace_document_chunks
    from backend.services.knowledge_storage import persist_knowledge_upload
    from backend.storage import get_storage

    upload = persist_knowledge_upload(
        storage=get_storage(),
        owner_id=owner_id,
        original_name=filename,
        content=content,
        content_type="text/plain",
        retention_policy="retained",
    )
    replace_document_chunks(upload.document_id, chunks)
    document = get_document(upload.document_id, owner_id)
    assert document is not None
    return document


def _docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("DOCX knowledge content")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _pptx_bytes() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box.text = "PPTX knowledge content"
    stream = BytesIO()
    presentation.save(stream)
    return stream.getvalue()


def _pdf_bytes() -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "PDF knowledge content")
    content = document.tobytes()
    document.close()
    return content


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "filename,body,expected",
    [
        ("notes.txt", b"TXT knowledge content", "TXT knowledge content"),
        ("notes.md", b"# Markdown knowledge content", "Markdown knowledge content"),
        ("notes.pdf", _pdf_bytes(), "PDF knowledge content"),
        ("notes.docx", _docx_bytes(), "DOCX knowledge content"),
        ("notes.pptx", _pptx_bytes(), "PPTX knowledge content"),
    ],
    ids=["txt", "md", "pdf", "docx", "pptx"],
)
async def test_extract_text_supports_personal_knowledge_formats(filename, body, expected):
    from backend.rag.chunker import extract_text

    text = await extract_text(filename, body)
    assert expected in text


@pytest.mark.asyncio
async def test_extract_text_rejects_unsupported_personal_knowledge_format():
    from fastapi import HTTPException
    from backend.rag.chunker import extract_text

    with pytest.raises(HTTPException) as exc:
        await extract_text("notes.png", b"not supported")
    assert exc.value.status_code == 400


@pytest.mark.asyncio
async def test_task_only_parse_failure_durably_enqueues_cleanup(monkeypatch):
    from sqlalchemy import select

    from backend.db.models import (
        AssignmentRecord,
        CourseRecord,
        KnowledgeStorageRecord,
        UserRecord,
    )
    from backend.db.session import session_scope
    from backend.domain.knowledge_storage import (
        KNOWLEDGE_CLEANUP_TASK_ATTACH_FAILED,
    )
    from backend.knowledge.service import ingest_document

    owner_id = "task-parse-failure-owner"
    assignment_id = "task-parse-failure-assignment"
    with session_scope() as session:
        session.add(UserRecord(
            id=owner_id,
            username=owner_id,
            password_hash="hash",
            role="teacher",
            is_active=True,
        ))
        session.flush()
        session.add(CourseRecord(
            id="task-parse-failure-course",
            name="Course",
            teacher_id=owner_id,
        ))
        session.flush()
        session.add(AssignmentRecord(
            id=assignment_id,
            course_id="task-parse-failure-course",
            teacher_id=owner_id,
            name="Assignment",
            status="draft",
            version=1,
        ))

    async def fail_parse(_filename: str, _content: bytes) -> str:
        raise ValueError("synthetic parse failure")

    monkeypatch.setattr("backend.knowledge.service.extract_text", fail_parse)
    body = b"task-only bytes whose parser fails"
    with pytest.raises(ValueError, match="synthetic parse failure"):
        await ingest_document(
            owner_id=owner_id,
            original_name="broken.txt",
            content=body,
            content_type="text/plain",
            retention_policy="task_only",
            origin_assignment_id=assignment_id,
        )

    with session_scope() as session:
        ledger = session.scalar(select(KnowledgeStorageRecord).where(
            KnowledgeStorageRecord.owner_id == owner_id
        ))
        assert ledger is not None
        assert ledger.state == "cleanup_pending"
        assert ledger.cleanup_reason == KNOWLEDGE_CLEANUP_TASK_ATTACH_FAILED
        assert ledger.size_bytes == len(body)


@pytest.mark.asyncio
async def test_task_only_parse_cleanup_enqueue_failure_does_not_mask_parser_error(
    monkeypatch,
):
    from sqlalchemy import select

    from backend.db.models import (
        AssignmentRecord,
        CourseRecord,
        KnowledgeStorageRecord,
        UserRecord,
    )
    from backend.db.session import session_scope
    from backend.knowledge.service import ingest_document

    owner_id = "task-parse-cleanup-failure-owner"
    assignment_id = "task-parse-cleanup-failure-assignment"
    with session_scope() as session:
        session.add(UserRecord(
            id=owner_id,
            username=owner_id,
            password_hash="hash",
            role="teacher",
            is_active=True,
        ))
        session.flush()
        session.add(CourseRecord(
            id="task-parse-cleanup-failure-course",
            name="Course",
            teacher_id=owner_id,
        ))
        session.flush()
        session.add(AssignmentRecord(
            id=assignment_id,
            course_id="task-parse-cleanup-failure-course",
            teacher_id=owner_id,
            name="Assignment",
            status="draft",
            version=1,
        ))

    async def fail_parse(_filename: str, _content: bytes) -> str:
        raise ValueError("primary parser failure")

    def fail_cleanup(*_args, **_kwargs):
        raise RuntimeError("secondary cleanup enqueue failure")

    monkeypatch.setattr("backend.knowledge.service.extract_text", fail_parse)
    monkeypatch.setattr(
        "backend.db.knowledge_storage_repository.request_task_only_cleanup_if_unreferenced",
        fail_cleanup,
    )
    with pytest.raises(ValueError, match="primary parser failure"):
        await ingest_document(
            owner_id=owner_id,
            original_name="broken.txt",
            content=b"task-only bytes with two failures",
            content_type="text/plain",
            retention_policy="task_only",
            origin_assignment_id=assignment_id,
        )

    with session_scope() as session:
        ledger = session.scalar(select(KnowledgeStorageRecord).where(
            KnowledgeStorageRecord.owner_id == owner_id
        ))
        assert ledger is not None
        assert ledger.state == "available"
        assert ledger.unattached_expires_at is not None


def test_knowledge_repository_is_owner_scoped_and_persists_assignment_selection():
    from backend.db.knowledge_repository import (
        get_document,
        list_documents,
        list_selected_documents,
        set_task_documents,
    )
    from backend.db.models import AssignmentRecord, CourseRecord, UserRecord
    from backend.db.session import session_scope

    with session_scope() as session:
        session.add_all([
            UserRecord(id="owner-a", username="owner-a", password_hash="hash", role="teacher",
                       is_active=True),
            UserRecord(id="owner-b", username="owner-b", password_hash="hash", role="teacher",
                       is_active=True),
        ])
        session.flush()
        # A course + assignment owned by owner-a; selection is assignment-scoped.
        session.add(CourseRecord(id="course-a", name="C", teacher_id="owner-a"))
        session.flush()
        session.add(AssignmentRecord(id="asg-a", course_id="course-a", teacher_id="owner-a",
                                     name="A", status="draft", version=1))

    document = _persist_ready_document(
        owner_id="owner-a",
        filename="notes.txt",
        content=b"quadratic formula and factorization",
        chunks=["quadratic formula", "factorization"],
    )

    assert get_document(document.id, "owner-a").status == "ready"
    assert get_document(document.id, "owner-b") is None
    assert [item.id for item in list_documents("owner-a")] == [document.id]
    assert list_documents("owner-b") == []

    set_task_documents(assignment_id="asg-a", owner_id="owner-a", document_ids=[document.id])
    assert [item.id for item in list_selected_documents("asg-a", "owner-a")] == [document.id]
    # Cross-owner selection resolves nothing (assignment owner predicate).
    assert list_selected_documents("asg-a", "owner-b") == []


@pytest.mark.asyncio
async def test_persistent_bm25_retriever_only_uses_selected_documents():
    from backend.db.knowledge_repository import set_task_documents
    from backend.db.models import AssignmentRecord, CourseRecord, UserRecord
    from backend.db.session import session_scope
    from backend.knowledge.retriever import PersistentKnowledgeRetriever

    with session_scope() as session:
        session.add(UserRecord(id="retriever-owner", username="retriever-owner", password_hash="hash",
                               role="teacher", is_active=True))
        session.flush()
        session.add(CourseRecord(id="rc", name="C", teacher_id="retriever-owner"))
        session.flush()
        session.add(AssignmentRecord(id="retriever-asg", course_id="rc", teacher_id="retriever-owner",
                                     name="A", status="draft", version=1))

    selected = _persist_ready_document(
        owner_id="retriever-owner",
        filename="selected.txt",
        content=b"Newton mechanics acceleration force",
        chunks=["Newton mechanics acceleration force"],
    )
    hidden = _persist_ready_document(
        owner_id="retriever-owner",
        filename="hidden.txt",
        content=b"Newton secret unselected sentence",
        chunks=["Newton secret unselected sentence"],
    )
    set_task_documents(assignment_id="retriever-asg", owner_id="retriever-owner", document_ids=[selected.id])

    # Scope is now the assignment id; the retriever resolves the teacher owner
    # through the assignment and only ranks the selected document's chunks.
    results = await PersistentKnowledgeRetriever().retrieve("Newton force", k=5, scope="retriever-asg")

    assert results
    assert results[0].source == "selected.txt"
    assert all("secret" not in item.content for item in results)


@pytest.mark.asyncio
async def test_grading_run_retriever_uses_frozen_owner_scoped_selection():
    from backend.db import grading_repository, workflow_repository
    from backend.db.knowledge_repository import (
        delete_document,
        set_task_documents,
    )
    from backend.db.models import AssignmentRecord, CourseRecord, UserRecord
    from backend.db.session import session_scope
    from backend.knowledge.retriever import PersistentKnowledgeRetriever
    from backend.domain.errors import InvalidTransition

    owner_id = "frozen-kb-owner"
    assignment_id = "frozen-kb-assignment"
    with session_scope() as session:
        session.add(UserRecord(
            id=owner_id, username=owner_id, password_hash="hash",
            role="teacher", is_active=True,
        ))
        session.flush()
        session.add(CourseRecord(id="frozen-kb-course", name="C", teacher_id=owner_id))
        session.flush()
        session.add(AssignmentRecord(
            id=assignment_id, course_id="frozen-kb-course", teacher_id=owner_id,
            name="A", status="draft", version=1,
        ))

    frozen = _persist_ready_document(
        owner_id=owner_id,
        filename="frozen.txt",
        content=b"Newton frozen knowledge force",
        chunks=["Newton frozen knowledge force"],
    )
    later = _persist_ready_document(
        owner_id=owner_id,
        filename="later.txt",
        content=b"Newton later selection force",
        chunks=["Newton later selection force"],
    )
    set_task_documents(
        assignment_id=assignment_id, owner_id=owner_id,
        document_ids=[frozen.id],
    )
    workflow = workflow_repository.ensure_workflow(
        assignment_id=assignment_id, owner_id=owner_id
    )
    run = grading_repository.create_run_bundle(
        assignment_id,
        teacher_id=owner_id,
        revision_ids=[],
        setup={"knowledge_scope": "all_task_docs"},
        setup_fingerprint="frozen-setup",
        input_manifest={"knowledge_document_ids": [frozen.id]},
        workflow_expected_revision=workflow.workflow_revision,
    )

    # Editing the assignment selection after the run is queued must not widen
    # or replace the knowledge visible to that grading run.
    set_task_documents(
        assignment_id=assignment_id, owner_id=owner_id,
        document_ids=[later.id],
    )
    results = await PersistentKnowledgeRetriever().retrieve(
        "Newton force", k=5, scope=f"grading-run:{run.id}",
    )

    assert results
    assert {item.source for item in results} == {"frozen.txt"}
    assert all("later" not in item.content for item in results)

    with pytest.raises(InvalidTransition) as exc:
        delete_document(frozen.id, owner_id)
    assert exc.value.code == "knowledge_document_in_active_grading_run"
    grading_repository.cancel(run.id, teacher_id=owner_id)
    assert delete_document(frozen.id, owner_id) is not None


def test_personal_knowledge_api_upload_list_download_and_delete():
    from sqlalchemy import select

    from backend.auth import create_token
    from backend.db.knowledge_repository import (
        list_selected_documents,
        set_task_documents,
    )
    from backend.db.models import AssignmentKnowledgeDocumentRecord, UserRecord
    from backend.db.session import session_scope
    from backend.main import app
    from backend.services import task_facade

    with session_scope() as session:
        session.add(UserRecord(id="demo_api-kb-owner", username="api-kb-owner", password_hash="hash",
                               role="teacher", is_active=True))

    client = TestClient(app)
    headers = {"Authorization": f"Bearer {create_token('demo_api-kb-owner', 'teacher')}"}
    uploaded = client.post("/knowledge/documents", headers=headers,
                           files={"file": ("api-notes.txt", b"persistent API knowledge", "text/plain")})
    assert uploaded.status_code == 201, uploaded.text
    document = uploaded.json()
    assert document["status"] == "ready"
    task = task_facade.create_task(
        owner_id="demo_api-kb-owner",
        name="Personal knowledge deletion",
        semester_id=None,
        course_id=None,
        idempotency_key="personal-kb-delete-detach",
    )
    set_task_documents(
        assignment_id=task["task_id"],
        owner_id="demo_api-kb-owner",
        document_ids=[document["id"]],
    )

    listed = client.get("/knowledge/documents", headers=headers)
    assert listed.status_code == 200
    assert [item["id"] for item in listed.json()["documents"]] == [document["id"]]

    downloaded = client.get(f"/knowledge/documents/{document['id']}/download", headers=headers)
    assert downloaded.status_code == 200
    assert downloaded.content == b"persistent API knowledge"

    assert client.get(f"/knowledge/documents/{document['id']}").status_code == 401
    removed = client.delete(f"/knowledge/documents/{document['id']}", headers=headers)
    assert removed.status_code == 202
    assert removed.json()["status"] == "deletion_pending"
    assert removed.json()["cleanup_operation_id"]
    assert client.get(f"/knowledge/documents/{document['id']}", headers=headers).status_code == 404
    assert list_selected_documents(task["task_id"], "demo_api-kb-owner") == []
    with session_scope() as session:
        assert session.scalar(select(AssignmentKnowledgeDocumentRecord).where(
            AssignmentKnowledgeDocumentRecord.assignment_id == task["task_id"],
            AssignmentKnowledgeDocumentRecord.document_id == document["id"],
        )) is None


def test_knowledge_storage_usage_api_reports_pending_retrying_and_reserved_bytes():
    import hashlib

    from backend.auth import create_token
    from backend.db.knowledge_storage_repository import (
        claim_next_cleanup,
        finish_cleanup,
        request_document_cleanup,
        reserve_upload,
    )
    from backend.db.models import UserRecord
    from backend.db.session import session_scope
    from backend.main import app

    owner_id = "demo_api-kb-usage-owner"
    with session_scope() as session:
        session.add(UserRecord(
            id=owner_id,
            username="api-kb-usage-owner",
            password_hash="hash",
            role="teacher",
            is_active=True,
        ))

    available_body = b"available knowledge"
    retrying_body = b"retrying knowledge"
    pending_body = b"pending knowledge"
    reserved_body = b"reserved knowledge"
    _persist_ready_document(
        owner_id=owner_id,
        filename="available.txt",
        content=available_body,
        chunks=["available knowledge"],
    )
    retrying = _persist_ready_document(
        owner_id=owner_id,
        filename="retrying.txt",
        content=retrying_body,
        chunks=["retrying knowledge"],
    )
    request_document_cleanup(retrying.id, owner_id)
    claim = claim_next_cleanup(worker_id="usage-test-worker")
    assert claim is not None
    assert claim.document_id == retrying.id
    finish_cleanup(
        claim,
        deleted=False,
        error_code="knowledge_storage_delete_failed",
    )
    pending = _persist_ready_document(
        owner_id=owner_id,
        filename="pending.txt",
        content=pending_body,
        chunks=["pending knowledge"],
    )
    request_document_cleanup(pending.id, owner_id)
    reserve_upload(
        owner_id=owner_id,
        original_name="reserved.txt",
        size_bytes=len(reserved_body),
        sha256=hashlib.sha256(reserved_body).hexdigest(),
        content_type="text/plain",
        retention_policy="retained",
    )

    client = TestClient(app)
    headers = {"Authorization": f"Bearer {create_token(owner_id, 'teacher')}"}
    response = client.get("/knowledge/storage/usage", headers=headers)
    assert response.status_code == 200, response.text
    usage = response.json()
    assert usage["used_bytes"] == sum(map(len, (
        available_body, retrying_body, pending_body, reserved_body,
    )))
    assert usage["available_document_bytes"] == len(available_body)
    assert usage["cleanup_pending_bytes"] == len(retrying_body) + len(pending_body)
    assert usage["retrying_cleanup_bytes"] == len(retrying_body)
    assert usage["reserved_bytes"] == len(reserved_body)
    assert usage["cleanup_pending_count"] == 2
    assert usage["retrying_cleanup_count"] == 1
    assert usage["reserved_count"] == 1


def test_assignment_knowledge_selection_caps_at_three_ready_documents():
    """A teacher may select at most three ready personal documents per assignment,
    and the selection survives repository recreation (DB is the source of truth)."""
    from backend.db.knowledge_repository import (
        list_selected_documents, set_task_documents,
    )
    from backend.db.models import AssignmentRecord, CourseRecord, UserRecord
    from backend.db.session import session_scope

    with session_scope() as session:
        session.add(UserRecord(id="cap-owner", username="cap-owner", password_hash="hash", role="teacher", is_active=True))
        session.flush()
        session.add(CourseRecord(id="cap-c", name="C", teacher_id="cap-owner"))
        session.flush()
        session.add(AssignmentRecord(id="cap-asg", course_id="cap-c", teacher_id="cap-owner", name="A",
                                     status="draft", version=1))

    doc_ids = []
    for i in range(4):
        doc = _persist_ready_document(
            owner_id="cap-owner",
            filename=f"d{i}.txt",
            content=f"unique content {i}".encode(),
            chunks=[f"chunk content {i}"],
        )
        doc_ids.append(doc.id)

    # Selecting four must be rejected.
    with pytest.raises(ValueError):
        set_task_documents(assignment_id="cap-asg", owner_id="cap-owner", document_ids=doc_ids)

    # Three is the cap and persists across a fresh repository call.
    set_task_documents(assignment_id="cap-asg", owner_id="cap-owner", document_ids=doc_ids[:3])
    assert [d.id for d in list_selected_documents("cap-asg", "cap-owner")] == doc_ids[:3]
